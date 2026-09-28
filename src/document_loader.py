import os
from docx import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document as LCDocument

DATA_DIR = "./data"

def load_docx(file_path):
    doc = Document(file_path)
    full_text = []
    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text.strip())
    return "\n".join(full_text)

def load_and_split_documents():
    documents = []
    if not os.path.exists(DATA_DIR):
        print(f"⚠️ Thư mục '{DATA_DIR}' không tồn tại.")
        return []

    for file_name in os.listdir(DATA_DIR):
        if file_name.startswith("~$") or file_name.startswith("."):
            continue

        file_path = os.path.join(DATA_DIR, file_name)
        if file_name.endswith(".docx"):
            text = load_docx(file_path)
            if text:
                documents.append(LCDocument(page_content=text, metadata={"source": file_name}))

    if not documents:
        print("⚠️ Không tìm thấy file hợp lệ trong thư mục data/")
        return []

    # Chunk_size 200 phù hợp hơn cho các đoạn văn bản ngắn
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=200,
        chunk_overlap=20,
        separators=["\n\n", "\n", " ", ""]
    )
    
    chunks = text_splitter.split_documents(documents)
    print(f"📄 Đã đọc và cắt thành {len(chunks)} đoạn văn bản (chunks).")
    return chunks